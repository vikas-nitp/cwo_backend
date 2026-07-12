from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol

import yaml


@dataclass(frozen=True)
class SourceSpec:
    path: Path
    format: str
    platform_id: str
    platform_name: str
    sheet: str | None = None


@dataclass(frozen=True)
class SourceRecord:
    data: dict[str, Any]
    source: str
    location: int
    sheet: str | None = None


class OfferSource(Protocol):
    def read(self, spec: SourceSpec, display_path: str) -> list[SourceRecord]: ...


class CsvOfferSource:
    def read(self, spec: SourceSpec, display_path: str) -> list[SourceRecord]:
        with spec.path.open(newline="", encoding="utf-8-sig") as handle:
            return [
                SourceRecord(dict(row), display_path, index)
                for index, row in enumerate(csv.DictReader(handle), 2)
            ]


class ExcelOfferSource:
    def read(self, spec: SourceSpec, display_path: str) -> list[SourceRecord]:
        from openpyxl import load_workbook

        workbook = load_workbook(spec.path, read_only=True, data_only=True)
        sheet_name = spec.sheet or "offers"
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"sheet '{sheet_name}' not found in {display_path}")
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [
            str(value).strip() if value is not None else "" for value in next(rows, ())
        ]
        return [
            SourceRecord(
                dict(zip(headers, values, strict=False)),
                display_path,
                index,
                sheet_name,
            )
            for index, values in enumerate(rows, 2)
        ]


class JsonOfferSource:
    def read(self, spec: SourceSpec, display_path: str) -> list[SourceRecord]:
        payload = json.loads(spec.path.read_text(encoding="utf-8-sig"))
        if isinstance(payload, dict) and set(payload) == {"offers"}:
            payload = payload["offers"]
        if not isinstance(payload, list) or not all(
            isinstance(item, dict) for item in payload
        ):
            raise ValueError(
                f"unsupported JSON shape in {display_path}; expected an array or {{'offers': [...]}}"
            )
        return [
            SourceRecord(dict(item), display_path, index)
            for index, item in enumerate(payload)
        ]


READERS: dict[str, OfferSource] = {
    "csv": CsvOfferSource(),
    "xlsx": ExcelOfferSource(),
    "json": JsonOfferSource(),
}


def load_catalogue(path: Path) -> list[SourceSpec]:
    payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict) or not isinstance(payload.get("sources"), list):
        raise ValueError("catalogue.yml must contain a sources list")
    specs: list[SourceSpec] = []
    for index, raw in enumerate(payload["sources"]):
        if not isinstance(raw, dict):
            raise ValueError(f"catalogue source {index} must be an object")
        format_name = str(raw.get("format", "")).lower()
        if format_name not in READERS:
            raise ValueError(f"unsupported source format: {format_name}")
        relative = Path(str(raw.get("path", "")))
        specs.append(
            SourceSpec(
                path=(path.parent / relative).resolve(),
                format=format_name,
                platform_id=str(raw.get("platform_id", "")).upper(),
                platform_name=str(raw.get("platform_name", "")),
                sheet=str(raw["sheet"]) if raw.get("sheet") else None,
            )
        )
    return specs


def read_source(spec: SourceSpec, catalogue_dir: Path) -> list[SourceRecord]:
    try:
        display_path = spec.path.relative_to(catalogue_dir.resolve()).as_posix()
    except ValueError:
        display_path = spec.path.name
    return READERS[spec.format].read(spec, display_path)
